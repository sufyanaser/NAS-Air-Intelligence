"""Phase 6 — Excel export: an 8-sheet study workbook built with openpyxl.

SQLite remains the source of truth; this module only renders what ``report.py`` and
``programming.py`` already computed from it.

Hard reporting gates:
- Sheets NEVER silently appear empty; UNAVAILABLE / BLOCKED / INSUFFICIENT_SAMPLE
  are displayed explicitly.
- The summary sheet separates CAPTURE QUALITY, ANALYSIS VALIDITY, and DECISION READINESS.
- Confidence scales are normalized to 0-1.
- All timestamps include explicit timezone context.
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import datetime, timedelta, tzinfo
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .. import __version__
from ..db import Database
from ..util import parse_iso
from .programming import build_programming_analysis
from .report import build_agent_report
from .timeline import station_timezone, timezone_metadata
from .validity import DecisionReadiness

SHEET_NAMES = [
    "01_Summary",
    "02_Timeline",
    "03_Content_Blocks",
    "04_Programs_Clock",
    "05_Dayparts",
    "06_Recurrent_Elements",
    "07_Programming_Insights",
    "08_Incidents_Technical",
]

_ARABIC_RANGES = (
    (0x0600, 0x06FF),
    (0x0750, 0x077F),
    (0x08A0, 0x08FF),
    (0xFB50, 0xFDFF),
    (0xFE70, 0xFEFF),
)
_INVALID_SHEET_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def _contains_arabic(text: str) -> bool:
    return any(any(lo <= ord(ch) <= hi for lo, hi in _ARABIC_RANGES) for ch in text)


def _write_cell(ws: Worksheet, row: int, col: int, value: Any, *, wrap: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and _contains_arabic(value):
        cell.alignment = Alignment(horizontal="right", wrap_text=wrap, readingOrder=2)
    elif wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _finalize_table(
    ws: Worksheet, headers: list[str], row_count: int, *, widths: list[int] | None = None
) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    last_row = max(row_count + 1, 2)
    ref = f"A1:{last_col}{last_row}"
    if row_count > 0:
        table_name = re.sub(r"[^A-Za-z0-9_]", "_", f"tbl_{ws.title}")
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = ref
    for col, width in enumerate(widths or [18] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _datetime_cell(iso_value: str | None, tz: tzinfo) -> datetime | None:
    if not iso_value:
        return None
    return parse_iso(iso_value).astimezone(tz).replace(tzinfo=None)


def _safe_filename_part(text: str) -> str:
    cleaned = _INVALID_SHEET_FILENAME_CHARS.sub("_", text).strip()
    return cleaned or "station"


def default_filename(station: str, started_at_local: datetime) -> str:
    return f"NAS-Air_{_safe_filename_part(station)}_{started_at_local:%Y-%m-%d}.xlsx"


# ------------------------------------------------------------------------------------- sheets
def _sheet_summary(
    ws: Worksheet,
    report: dict[str, Any],
    programming: dict[str, Any],
    events: list[dict[str, Any]],
) -> None:
    info = report["monitoring_information"]
    tz = station_timezone(info.get("timezone"))
    tz_meta = timezone_metadata(tz)
    model_device = Counter(
        (m.get("model"), m.get("device"), m.get("compute_type"))
        for e in events
        if (m := e.get("metadata") or {}).get("model")
    )
    if model_device:
        model, device, compute_type = model_device.most_common(1)[0][0]
    else:
        model, device, compute_type = None, None, None

    stream = info.get("stream") or {}
    status_block = programming.get("status") or {}
    cov = programming.get("coverage_metrics") or {}
    validity = programming.get("validity") or {}
    readiness = status_block.get("decision_readiness", DecisionReadiness.NOT_READY.value)

    # Styling helper
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    r = 1

    def _section_header(title: str) -> None:
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = header_font
        ws.cell(row=r, column=1).fill = header_fill
        ws.cell(row=r, column=2, value="").fill = header_fill
        r += 1

    def _add_row(label: str, val: Any, is_date: bool = False, wrap: bool = False) -> None:
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        _write_cell(ws, r, 2, val)
        if is_date and isinstance(val, datetime):
            ws.cell(row=r, column=2).number_format = "yyyy-mm-dd hh:mm:ss"
        if wrap:
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
        r += 1

    # SECTION 1: DECISION READINESS
    _section_header("DECISION READINESS")
    _add_row("Decision Readiness", readiness)
    if validity.get("reasons"):
        _add_row("Readiness Notes / Blocker", " | ".join(validity["reasons"]), wrap=True)
    if validity.get("limitations"):
        _add_row("Active Limitations", " | ".join(validity["limitations"]), wrap=True)
    r += 1

    # SECTION 2: CAPTURE QUALITY
    _section_header("CAPTURE QUALITY")
    _add_row("Capture Status", status_block.get("capture_status", "UNKNOWN"))
    _add_row("Station", info["station"])
    _add_row("Quality verdict", report["outcome"])
    _add_row("Official Page", info.get("source_page"))
    _add_row("Resolved stream", stream.get("resolved_stream_url") or info.get("stream_url"))
    _add_row("Session Start", _datetime_cell(info["started_at"], tz), is_date=True)
    _add_row("Session End", _datetime_cell(info["ended_at"], tz), is_date=True)
    _add_row("Requested Duration (s)", info["requested_seconds"])
    _add_row(
        "Captured Audio Duration (s)",
        cov.get("captured_seconds", report["capture_health"]["captured_seconds"]),
    )
    _add_row(
        "Captured Chunks Count", cov.get("chunk_count", report["capture_health"]["chunk_count"])
    )
    _add_row("Unique Covered Audio (s)", cov.get("unique_covered_seconds", 0.0))
    _add_row("Capture Gap Duration (s)", cov.get("gap_seconds", 0.0))
    _add_row("Chunk Overlap Duration (s)", cov.get("overlap_seconds", 0.0))
    _add_row("Timeline Coverage Ratio (%)", cov.get("coverage_percent", 0.0))
    _add_row("Incidents (Non-informational)", report["incidents"]["total"])
    r += 1

    # SECTION 3: ANALYSIS VALIDITY
    _section_header("ANALYSIS VALIDITY")
    _add_row("Processing Pipeline Status", status_block.get("processing_status", "UNKNOWN"))
    _add_row("Speech Transcription Status", status_block.get("transcription_status", "UNKNOWN"))
    _add_row(
        "Broadcast Classification Status", status_block.get("classification_status", "UNKNOWN")
    )
    _add_row(
        "Programming Analysis Status", status_block.get("programming_analysis_status", "UNKNOWN")
    )
    _add_row(
        "Mean Content Block Confidence",
        round(
            sum(b["confidence"] for b in programming["content_blocks"])
            / len(programming["content_blocks"]),
            3,
        )
        if programming.get("content_blocks")
        else "UNAVAILABLE",
    )
    r += 1

    # SECTION 4: ENVIRONMENT & METADATA
    _section_header("SYSTEM ENVIRONMENT & METADATA")
    _add_row("Schema Version", 1)
    _add_row("App Version", __version__)
    _add_row("Station Timezone", f"{tz_meta['timezone']} ({tz_meta['offset_iso']})")
    _add_row("Speech Engine Model", model or "n/a")
    _add_row("Compute Device", device or "n/a")
    _add_row("Compute Type", compute_type or "n/a")

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 60
    ws.print_area = f"A1:B{r - 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _segment_evidence(events_in_window: list[dict[str, Any]]) -> tuple[str | None, str]:
    if not events_in_window:
        return None, "no contributing event"
    texts = [str(e["text"]).strip() for e in events_in_window if e.get("text")]
    labels = sorted({e.get("label") for e in events_in_window if e.get("label")})
    transcript = " ".join(texts) if texts else None
    evidence = f"{len(events_in_window)} event(s)" + (f": {', '.join(labels)}" if labels else "")
    return transcript, evidence


def _sheet_timeline(
    ws: Worksheet,
    report: dict[str, Any],
    events: list[dict[str, Any]],
    chunks: list[dict[str, Any]],
) -> None:
    headers = ["Start", "End", "Duration (s)", "Category", "Confidence", "Transcript", "Evidence"]
    tz = station_timezone(report["monitoring_information"].get("timezone"))
    chunk_start = {c["id"]: parse_iso(c["started_at"]) for c in chunks}
    segments = report["timeline"]["segments"]

    if not segments:
        _write_cell(ws, 2, 1, "Status: UNAVAILABLE / EMPTY — No timeline segments recorded.")
        _finalize_table(ws, headers, 1, widths=[20, 20, 12, 16, 12, 60, 24])
        return

    for row, seg in enumerate(segments, start=2):
        start_dt = parse_iso(seg["start"])
        end_dt = parse_iso(seg["end"])
        window_events = [
            e
            for e in events
            if e.get("chunk_id") in chunk_start
            and chunk_start[e["chunk_id"]] + timedelta(seconds=float(e["start_offset"])) < end_dt
            and chunk_start[e["chunk_id"]] + timedelta(seconds=float(e["end_offset"])) > start_dt
        ]
        transcript, evidence = _segment_evidence(window_events)
        _write_cell(ws, row, 1, start_dt.astimezone(tz).replace(tzinfo=None))
        ws.cell(row=row, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 2, end_dt.astimezone(tz).replace(tzinfo=None))
        ws.cell(row=row, column=2).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 3, seg["duration_seconds"])
        _write_cell(ws, row, 4, seg["kind"])
        _write_cell(ws, row, 5, seg["mean_confidence"])
        _write_cell(ws, row, 6, transcript, wrap=True)
        _write_cell(ws, row, 7, evidence)
    _finalize_table(ws, headers, len(segments), widths=[20, 20, 12, 16, 12, 60, 24])


def _sheet_content_blocks(ws: Worksheet, programming: dict[str, Any], tz: tzinfo) -> None:
    headers = ["Start", "End", "Duration (s)", "Block Type", "Evidence", "Confidence", "Notes"]
    blocks = programming.get("content_blocks", [])
    status = programming.get("status", {}).get("programming_analysis_status")

    if not blocks:
        note = (
            f"Status: {status or 'UNAVAILABLE'} — Content blocks unavailable: "
            "transcription/classification dependencies failed."
        )
        _write_cell(ws, 2, 1, note)
        _finalize_table(ws, headers, 1, widths=[20, 20, 12, 18, 24, 12, 40])
        return

    for row, b in enumerate(blocks, start=2):
        _write_cell(ws, row, 1, parse_iso(b["start"]).astimezone(tz).replace(tzinfo=None))
        ws.cell(row=row, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 2, parse_iso(b["end"]).astimezone(tz).replace(tzinfo=None))
        ws.cell(row=row, column=2).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 3, b["duration_seconds"])
        _write_cell(ws, row, 4, b["block_type"])
        ev_cnt = len(b.get("evidence_event_ids", []))
        ev_label = f"{ev_cnt} event(s)" if ev_cnt else b.get("evidence", "OBSERVED")
        _write_cell(ws, row, 5, ev_label)
        _write_cell(ws, row, 6, b["confidence"])
        notes_val = b.get("notes") or ""
        if b.get("id") and not notes_val:
            notes_val = f"Block {b.get('id')}"
        _write_cell(ws, row, 7, notes_val, wrap=True)
    _finalize_table(ws, headers, len(blocks), widths=[20, 20, 12, 18, 24, 12, 40])


def _sheet_programs_clock(ws: Worksheet, programming: dict[str, Any]) -> None:
    headers = [
        "Type",
        "Position",
        "Detail",
        "Occurrences",
        "Mean Duration (s)",
        "Evidence",
        "Confidence",
        "Sufficiency Status",
    ]
    candidates = programming.get("program_candidates", [])
    patterns = programming.get("clock_patterns", [])
    status = programming.get("status", {}).get("programming_analysis_status")

    rows: list[tuple[Any, ...]] = []
    for c in candidates:
        rows.append(
            (
                "Program Candidate",
                c["bucket"],
                c["block_type"],
                c["occurrences"],
                c["mean_duration_seconds"],
                "; ".join(c["evidence"][:3]),
                c["confidence"],
                c.get("sufficiency_status", "LIMITED"),
            )
        )
    for p in patterns:
        rows.append(
            (
                "Clock Pattern",
                p["pattern_type"],
                p["description"],
                p["occurrences"],
                None,
                "; ".join(p["evidence"][:3]),
                p["confidence"],
                p.get("sufficiency_status", "LIMITED"),
            )
        )

    if not rows:
        reason = (
            f"Status: {status or 'UNAVAILABLE'} — Clock patterns and candidates not available; "
            "insufficient monitoring sample (< 1h) or transcription dependency failed."
        )
        rows.append(("Note", "n/a", reason, 0, None, "n/a", 0.0, status or "UNAVAILABLE"))

    for row, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            _write_cell(ws, row, col, value, wrap=(col in (3, 6)))
    _finalize_table(ws, headers, len(rows), widths=[18, 14, 44, 12, 16, 40, 12, 18])


def _sheet_dayparts(ws: Worksheet, programming: dict[str, Any]) -> None:
    headers = [
        "Daypart",
        "Window",
        "Monitored (s)",
        "Speech (s)",
        "Unknown (s)",
        "Silence (s)",
        "Avg Block Duration (s)",
        "Presenter Return Interval (s)",
        "Recurrent Elements",
        "Program Candidates",
        "Observations",
        "Confidence",
        "Sufficiency Status",
    ]
    dayparts = programming.get("dayparts", [])
    status = programming.get("status", {}).get("programming_analysis_status")

    if not dayparts:
        _write_cell(
            ws, 2, 1, f"Status: {status or 'UNAVAILABLE'} — Daypart analysis not generated."
        )
        _finalize_table(ws, headers, 1, widths=[12, 14, 14, 12, 14, 12, 18, 22, 14, 14, 50, 12, 18])
        return

    for row, d in enumerate(dayparts, start=2):
        speech_val = d["speech_seconds"] if d["speech_seconds"] is not None else "UNAVAILABLE"
        unknown_val = d["unknown_seconds"] if d["unknown_seconds"] is not None else "UNAVAILABLE"
        values = [
            d["daypart"],
            d["window"],
            d["monitored_seconds"],
            speech_val,
            unknown_val,
            d["silence_seconds"],
            d["avg_block_duration_seconds"],
            d["presenter_return_interval_seconds"]
            if d["presenter_return_interval_seconds"] is not None
            else "UNAVAILABLE",
            d["recurrent_element_count"],
            d["program_candidate_count"],
            " | ".join(o["text"] for o in d["observations"])
            or "Analysis BLOCKED or INSUFFICIENT_SAMPLE",
            d["confidence"],
            d.get("sufficiency_status", "INSUFFICIENT_SAMPLE"),
        ]
        for col, value in enumerate(values, start=1):
            _write_cell(ws, row, col, value, wrap=(col == 11))
    _finalize_table(
        ws,
        headers,
        len(dayparts),
        widths=[12, 14, 14, 14, 14, 12, 18, 22, 14, 14, 50, 12, 18],
    )


def _sheet_recurrent_elements(ws: Worksheet, programming: dict[str, Any], tz: tzinfo) -> None:
    headers = [
        "Candidate",
        "Occurrences",
        "First Seen",
        "Last Seen",
        "Evidence",
        "Confidence",
        "Status",
    ]
    elements = programming.get("recurrent_elements", [])

    if not elements:
        _write_cell(ws, 2, 1, "Status: NONE_DETECTED — No repeated acoustic fingerprints observed.")
        _finalize_table(ws, headers, 1, widths=[24, 12, 20, 20, 28, 12, 14])
        return

    for row, e in enumerate(elements, start=2):
        starts = sorted(parse_iso(s) for s in e["chunk_starts"])
        confidence = round(min(0.5, 0.15 * e["occurrences"]), 3)
        _write_cell(ws, row, 1, e["label"])
        _write_cell(ws, row, 2, e["occurrences"])
        _write_cell(ws, row, 3, starts[0].astimezone(tz).replace(tzinfo=None) if starts else None)
        ws.cell(row=row, column=3).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 4, starts[-1].astimezone(tz).replace(tzinfo=None) if starts else None)
        ws.cell(row=row, column=4).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_cell(ws, row, 5, f"fingerprint {e['fingerprint_prefix']}…")
        _write_cell(ws, row, 6, confidence)
        _write_cell(ws, row, 7, "NOT VERIFIED")
    _finalize_table(ws, headers, len(elements), widths=[24, 12, 20, 20, 28, 12, 14])


def _sheet_programming_insights(ws: Worksheet, programming: dict[str, Any]) -> None:
    headers = [
        "Type",
        "Observation",
        "Evidence",
        "Occurrences",
        "Confidence",
        "NAS FM Planning Input",
    ]
    insights = programming.get("insights", [])

    if not insights:
        _write_cell(
            ws,
            2,
            1,
            (
                "Status: BLOCKED — No programming insights or recommendations "
                "can be derived without verified evidence."
            ),
        )
        _finalize_table(ws, headers, 1, widths=[18, 48, 36, 12, 12, 48])
        return

    for row, i in enumerate(insights, start=2):
        _write_cell(ws, row, 1, i.get("type") or i.get("classification"))
        _write_cell(ws, row, 2, i["observation"], wrap=True)
        ev_text = i["evidence"]
        if i.get("confidence_basis"):
            ev_text = f"{ev_text} (Basis: {i['confidence_basis']})"
        _write_cell(ws, row, 3, ev_text, wrap=True)
        _write_cell(ws, row, 4, i["occurrences"])
        _write_cell(ws, row, 5, i["confidence"])
        _write_cell(ws, row, 6, i.get("nas_fm_planning_input") or "None", wrap=True)
    _finalize_table(ws, headers, len(insights), widths=[18, 48, 36, 12, 12, 48])


def _sheet_incidents_technical(ws: Worksheet, report: dict[str, Any]) -> None:
    items = report["incidents"]["items"]
    headers = ["Time", "Type", "Severity", "Duration", "Details"]
    if not items:
        _write_cell(ws, 2, 1, "No incidents recorded.")
        row_count = 1
    else:
        for row, i in enumerate(items, start=2):
            _write_cell(ws, row, 1, i["occurred_at"])
            _write_cell(ws, row, 2, i["kind"])
            _write_cell(ws, row, 3, i["severity"])
            _write_cell(ws, row, 4, "n/a")
            _write_cell(ws, row, 5, i["details"], wrap=True)
        row_count = len(items)
    _finalize_table(ws, headers, row_count, widths=[22, 18, 12, 10, 60])

    tech_header_row = row_count + 3
    ws.cell(row=tech_header_row, column=1, value="Technical / Quality Gates").font = Font(bold=True)
    gate_headers = ["Gate", "Status", "Detail"]
    for col, h in enumerate(gate_headers, start=1):
        ws.cell(row=tech_header_row + 1, column=col, value=h).font = Font(bold=True)
    for row, gate in enumerate(report["capture_health"]["gates"], start=tech_header_row + 2):
        _write_cell(ws, row, 1, gate["gate"])
        _write_cell(ws, row, 2, gate["status"])
        _write_cell(ws, row, 3, gate["detail"], wrap=True)
    for col, width in enumerate([22, 18, 12, 10, 60], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


# ------------------------------------------------------------------------------------- top level
def write_workbook(
    db: Database,
    session_id: str,
    output_dir: str | Path,
    *,
    requested_seconds: float,
    run: dict[str, Any] | None = None,
    filename: str | None = None,
) -> Path:
    """Builds the 8-sheet study workbook for one session and returns its path."""
    report = build_agent_report(db, session_id, requested_seconds=requested_seconds, run=run)
    programming = build_programming_analysis(db, session_id)
    events = db.events(session_id)
    chunks = db.chunks(session_id)
    tz = station_timezone(report["monitoring_information"].get("timezone"))

    wb = Workbook()
    wb.remove(wb.active)
    s01 = wb.create_sheet(SHEET_NAMES[0])
    s02 = wb.create_sheet(SHEET_NAMES[1])
    s03 = wb.create_sheet(SHEET_NAMES[2])
    s04 = wb.create_sheet(SHEET_NAMES[3])
    s05 = wb.create_sheet(SHEET_NAMES[4])
    s06 = wb.create_sheet(SHEET_NAMES[5])
    s07 = wb.create_sheet(SHEET_NAMES[6])
    s08 = wb.create_sheet(SHEET_NAMES[7])

    _sheet_summary(s01, report, programming, events)
    _sheet_timeline(s02, report, events, chunks)
    _sheet_content_blocks(s03, programming, tz)
    _sheet_programs_clock(s04, programming)
    _sheet_dayparts(s05, programming)
    _sheet_recurrent_elements(s06, programming, tz)
    _sheet_programming_insights(s07, programming)
    _sheet_incidents_technical(s08, report)

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    started_local = parse_iso(report["monitoring_information"]["started_at"]).astimezone(tz)
    name = filename or default_filename(report["monitoring_information"]["station"], started_local)
    path = out / name
    wb.save(path)
    return path
