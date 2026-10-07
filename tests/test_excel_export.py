"""Phase 6 Excel export: an 8-sheet workbook, Arabic preserved exactly, real date/time
values, no raw JSON, frozen headers + filter/table on every tabular sheet.
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from nas_air_intelligence.agent.excel_export import (
    SHEET_NAMES,
    _contains_arabic,
    default_filename,
    write_workbook,
)
from nas_air_intelligence.db import Database
from nas_air_intelligence.util import isoformat

ARABIC_TEXT = "مرحباً بكم في البرنامج"


def _build_session(
    tmp_path: Path, *, with_duplicate_fingerprint: bool = True
) -> tuple[Database, str]:
    db = Database(tmp_path / "study.db")
    station = db.upsert_station("محطة الاختبار", "https://s/live.mp3")
    session_id = db.create_session(station, 7200.0, 300, str(tmp_path / "out"))
    t0 = datetime(2026, 1, 5, 10, 0, 0, tzinfo=UTC)
    fp = "abc123fingerprintdata" * 3
    for h in range(2):
        for m in (0, 20, 40):
            start = t0 + timedelta(hours=h, minutes=m)
            chunk_id = db.add_chunk(
                session_id, f"/tmp/c{h}-{m}.mp3", isoformat(start),
                isoformat(start + timedelta(seconds=180)), 180.0, f"sha-{h}-{m}",
            )  # fmt: skip
            db.add_events([
                {"id": f"e-{h}-{m}", "session_id": session_id, "chunk_id": chunk_id,
                 "kind": "speech", "start_offset": 0.0, "end_offset": 150.0, "confidence": 0.8,
                 "text": ARABIC_TEXT, "fingerprint": fp if with_duplicate_fingerprint else None,
                 "metadata": {"model": "tiny", "device": "cpu"}},
            ])  # fmt: skip
    db.add_incident(session_id, "ffmpeg_exit", "returncode=1")
    db.finish_session(session_id, "completed")
    return db, session_id


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    db, session_id = _build_session(tmp_path)
    return write_workbook(db, session_id, tmp_path / "exports", requested_seconds=7200.0)


def test_eight_sheets_in_the_exact_spec_order(workbook_path: Path):
    wb = openpyxl.load_workbook(workbook_path)
    assert wb.sheetnames == SHEET_NAMES
    assert len(SHEET_NAMES) == 8


def test_filename_follows_the_recommended_pattern(workbook_path: Path):
    assert workbook_path.name.startswith("NAS-Air_")
    assert workbook_path.suffix == ".xlsx"
    expected = "NAS-Air_Al Nakhla FM_2026-03-01.xlsx"
    assert default_filename("Al Nakhla FM", datetime(2026, 3, 1)) == expected


def test_filename_sanitizes_filesystem_unsafe_characters():
    name = default_filename("A/B:C*D?", datetime(2026, 3, 1))
    assert not any(ch in name.removeprefix("NAS-Air_") for ch in '/:*?"<>|')


def test_summary_sheet_has_key_value_rows_not_a_raw_dump(workbook_path: Path):
    ws = openpyxl.load_workbook(workbook_path)["01_Summary"]
    labels = [row[0].value for row in ws.iter_rows(min_row=1, max_col=1)]
    assert "Station" in labels and "Quality verdict" in labels
    station_row = next(r for r in ws.iter_rows() if r[0].value == "Station")
    assert station_row[1].value == "محطة الاختبار"


def test_timeline_sheet_headers_dates_and_table(workbook_path: Path):
    ws = openpyxl.load_workbook(workbook_path)["02_Timeline"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Start", "End", "Duration (s)", "Category", "Confidence", "Transcript",
                        "Evidence"]  # fmt: skip
    assert ws.freeze_panes == "A2"
    assert ws.tables, "expected an Excel-native table on the timeline sheet"
    first_data_row = ws[2]
    assert isinstance(first_data_row[0].value, datetime)  # real date/time, not an ISO string


def test_arabic_transcript_round_trips_exactly_and_is_marked_rtl(workbook_path: Path):
    ws = openpyxl.load_workbook(workbook_path)["02_Timeline"]
    transcripts = [row[5].value for row in ws.iter_rows(min_row=2) if row[5].value]
    assert transcripts, "expected at least one transcript cell"
    assert all(t == ARABIC_TEXT for t in transcripts)
    arabic_cell = next(row[5] for row in ws.iter_rows(min_row=2) if row[5].value)
    assert arabic_cell.alignment.readingOrder == 2
    assert arabic_cell.alignment.horizontal == "right"


def test_technical_urls_and_timestamps_stay_left_to_right(workbook_path: Path):
    ws = openpyxl.load_workbook(workbook_path)["01_Summary"]
    stream_row = next(r for r in ws.iter_rows() if r[0].value == "Resolved stream")
    assert stream_row[1].alignment.readingOrder in (None, 0)


def test_content_blocks_and_programming_insights_sheets_have_headers_and_tables(
    workbook_path: Path,
):
    wb = openpyxl.load_workbook(workbook_path)
    blocks = wb["03_Content_Blocks"]
    assert [c.value for c in blocks[1]] == [
        "Start", "End", "Duration (s)", "Block Type", "Evidence", "Confidence", "Notes",
    ]  # fmt: skip
    insights = wb["07_Programming_Insights"]
    assert [c.value for c in insights[1]] == [
        "Type", "Observation", "Evidence", "Occurrences", "Confidence",
        "NAS FM Planning Input",
    ]  # fmt: skip
    assert insights.freeze_panes == "A2"


def test_recurrent_elements_sheet_uses_the_unidentified_candidate_label(tmp_path: Path):
    db, session_id = _build_session(tmp_path, with_duplicate_fingerprint=True)
    path = write_workbook(db, session_id, tmp_path / "exports", requested_seconds=7200.0)
    ws = openpyxl.load_workbook(path)["06_Recurrent_Elements"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows, "the fixture has a duplicated fingerprint and should produce a candidate row"
    assert rows[0][0] == "recurrent_audio_candidate"
    assert rows[0][6] == "NOT VERIFIED"


def test_no_cell_contains_a_raw_json_dump(workbook_path: Path):
    wb = openpyxl.load_workbook(workbook_path)
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    stripped = cell.value.strip()
                    assert not (stripped.startswith("{") or stripped.startswith("[")), (
                        f"{name}!{cell.coordinate} looks like a raw JSON/list dump: {cell.value!r}"
                    )


def test_incidents_sheet_lists_incidents_and_quality_gates(workbook_path: Path):
    ws = openpyxl.load_workbook(workbook_path)["08_Incidents_Technical"]
    assert [c.value for c in ws[1]] == ["Time", "Type", "Severity", "Duration", "Details"]
    kinds = [row[1].value for row in ws.iter_rows(min_row=2, max_row=ws.max_row) if row[1].value]
    assert "ffmpeg_exit" in kinds
    gate_header_row = next(
        r[0] for r in ws.iter_rows() if r[0].value == "Technical / Quality Gates"
    ).row
    gate_headers = [c.value for c in ws[gate_header_row + 1]]
    assert gate_headers[:3] == ["Gate", "Status", "Detail"]


def test_contains_arabic_helper():
    assert _contains_arabic("مرحبا")
    assert _contains_arabic("hello مرحبا")
    assert not _contains_arabic("hello https://example.com 2026-01-05")


def test_workbook_is_reproducible(tmp_path: Path):
    db, session_id = _build_session(tmp_path)
    first = write_workbook(db, session_id, tmp_path / "a", requested_seconds=7200.0)
    second = write_workbook(db, session_id, tmp_path / "b", requested_seconds=7200.0)
    wb1 = openpyxl.load_workbook(first)
    wb2 = openpyxl.load_workbook(second)
    for name in SHEET_NAMES:
        rows1 = list(wb1[name].iter_rows(values_only=True))
        rows2 = list(wb2[name].iter_rows(values_only=True))
        assert rows1 == rows2, f"{name} differs between identical runs"
